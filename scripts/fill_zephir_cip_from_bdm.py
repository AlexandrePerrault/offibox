#!/usr/bin/env python3
"""
Remplit la colonne B (CIP) du fichier Excel Zéphir à partir du BDM master.
Col A = Produit, Col B = CIP, Col C = URL.
Charge BDM_MASTER2026.csv (dénomination col 0, CIP13 col 2), fait la correspondance
par nom (recherche du nom produit dans la dénomination BDM).
Usage: python fill_zephir_cip_from_bdm.py [fichier.xlsx] [--out fichier_sortie.xlsx]
Si --out est omis, écrase le fichier source (fermez-le dans Excel avant).
"""
import csv
import io
import re
import sys
import unicodedata

try:
    import requests
    import openpyxl
except ImportError:
    print("Dépendances: requests, openpyxl")
    print("  python -m pip install requests openpyxl")
    sys.exit(1)

BDM_MASTER_URL = "https://raw.githubusercontent.com/AlexandrePerrault/offiboxdata/main/BDM_MASTER2026.csv"
DEFAULT_XLSX = "splf_zephir_videos.xlsx"


def normalize(s):
    """Minuscules, sans accents, alphanumeriques + espaces."""
    if not s:
        return ""
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9\s]", "", s.lower()).strip()


def extract_search_keys(produit):
    """Extrait des clés de recherche depuis le libellé produit Zéphir."""
    # Enlever (xxx) et ®, ®, etc.
    s = re.sub(r"\s*\([^)]*\)\s*", " ", produit)
    s = re.sub(r"[®™]", " ", s)
    s = s.strip()
    if not s:
        return []
    # Premier mot (souvent la marque: VENTOLINE, ANORO)
    words = re.split(r"[\s+,]+", s)
    keys = []
    for w in words:
        w = w.strip()
        if len(w) >= 2 and w.upper() not in ("MG", "ML", "G", "MG/ML", "ET", "DE", "DU"):
            keys.append(normalize(w))
    # Clé principale = premier mot significatif
    if keys:
        keys = [keys[0]] + [k for k in keys[1:] if len(k) >= 4][:2]
    return keys


def load_bdm_rows():
    """Télécharge et parse BDM_MASTER2026.csv. Retourne liste de (denomination, cip13)."""
    r = requests.get(BDM_MASTER_URL, timeout=60)
    r.raise_for_status()
    r.encoding = r.apparent_encoding or "utf-8"
    reader = csv.reader(io.StringIO(r.text), delimiter=";")
    next(reader, None)
    rows = []
    for row in reader:
        if len(row) < 3:
            continue
        denom = row[0].strip().strip('"')
        cip13 = re.sub(r"\D", "", row[2].strip())
        if len(cip13) != 13:
            continue
        rows.append((denom, cip13))
    return rows


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    out_path = None
    if "--out" in sys.argv:
        i = sys.argv.index("--out")
        if i + 1 < len(sys.argv):
            out_path = sys.argv[i + 1]
    xlsx_path = args[0] if args else DEFAULT_XLSX
    xlsx_path = str(xlsx_path)
    if out_path is None:
        out_path = xlsx_path
    try:
        wb = openpyxl.load_workbook(xlsx_path)
    except Exception as e:
        print(f"Fichier introuvable ou invalide: {xlsx_path}")
        print(e)
        sys.exit(1)

    ws = wb.active
    if ws.max_row < 2:
        print("Aucune donnée dans l'Excel.")
        sys.exit(0)

    print("Chargement du BDM master...")
    bdm_rows = load_bdm_rows()
    print(f"  {len(bdm_rows)} lignes BDM")

    bdm_flat = [(normalize(d), d, c) for d, c in bdm_rows]

    filled = 0
    for row in range(2, ws.max_row + 1):
        produit = ws.cell(row=row, column=1).value
        if not produit or not str(produit).strip():
            continue
        produit = str(produit).strip()
        keys = extract_search_keys(produit)
        if not keys:
            continue
        main_key = keys[0]
        cip_found = None
        for norm_denom, denom, cip in bdm_flat:
            if main_key in norm_denom:
                denom_upper = denom.upper()
                if any(
                    x in denom_upper
                    for x in (
                        "INHALATION",
                        "AÉROSOL",
                        "AEROSOL",
                        "DISKUS",
                        "ELLIPTA",
                        "TURBUHALER",
                        "AUTOHALER",
                        "NEXTHALER",
                        "BREEZHALER",
                        "RESPIMAT",
                        "SPRAY",
                        "INHAL",
                    )
                ):
                    cip_found = cip
                    break
        if cip_found is None:
            for norm_denom, denom, cip in bdm_flat:
                if main_key in norm_denom:
                    cip_found = cip
                    break
        if cip_found:
            ws.cell(row=row, column=2, value=cip_found)
            filled += 1

    wb.save(out_path)
    print(f"Colonne CIP remplie pour {filled} produits. Fichier enregistré: {out_path}")


if __name__ == "__main__":
    main()
