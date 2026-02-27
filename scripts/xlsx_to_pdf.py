#!/usr/bin/env python3
"""
Génère 2 PDF à partir des 2 feuilles d'un fichier Excel
(liste médicaments écrasables). Un PDF par feuille, format tableau lisible.
Usage: python xlsx_to_pdf.py [chemin.xlsx]
Sortie: 2 fichiers PDF dans le même dossier que le xlsx.
"""
from pathlib import Path
import sys

try:
    import openpyxl
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
except ImportError as e:
    print("Dépendances requises: openpyxl, reportlab")
    print("  python -m pip install openpyxl reportlab")
    sys.exit(1)


def safe_str(val):
    if val is None:
        return ""
    return str(val).strip()


def sheet_to_data(ws):
    """Lit toute la feuille en liste de listes (chaque cellule en str)."""
    rows = []
    for row in ws.iter_rows():
        rows.append([safe_str(c.value) for c in row])
    return rows


def build_pdf_from_sheet(sheet_name, rows, out_path):
    """Génère un PDF à partir des lignes (liste de listes)."""
    if not rows:
        return

    # Nombre de colonnes = max des largeurs de ligne
    ncols = max(len(r) for r in rows) if rows else 0
    if ncols == 0:
        return

    # Normaliser: toutes les lignes ont le même nombre de colonnes
    for r in rows:
        while len(r) < ncols:
            r.append("")

    # Choix orientation et taille police
    if ncols > 6:
        page_size = landscape(A4)
        font_size = 7
    else:
        page_size = A4
        font_size = 9

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=page_size,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=15 * mm,
        bottomMargin=12 * mm,
    )

    # Données pour la Table (chaque cellule = texte simple)
    table = Table(rows, repeatRows=1)
    table.setStyle(
        TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), font_size + 1),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E0E0E0")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F5")]),
        ])
    )

    doc.build([table])


def main():
    default_path = r"C:\Users\perra\Downloads\Liste médicaments écrasables_30012026.xlsx"
    xlsx_path = Path(sys.argv[1] if len(sys.argv) > 1 else default_path)

    if not xlsx_path.exists():
        print(f"Fichier introuvable: {xlsx_path}")
        print("Usage: python xlsx_to_pdf.py [chemin.xlsx]")
        sys.exit(1)

    out_dir = xlsx_path.parent
    base_name = xlsx_path.stem

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if len(wb.sheetnames) < 2:
        print("Le classeur doit contenir au moins 2 feuilles.")
        wb.close()
        sys.exit(1)

    for i, sheet_name in enumerate(wb.sheetnames[:2], start=1):
        ws = wb[sheet_name]
        rows = sheet_to_data(ws)
        # Nom de fichier sans caractères gênants
        safe_sheet = "".join(c if c.isalnum() or c in " _-" else "_" for c in sheet_name)
        out_name = f"{base_name}_{safe_sheet}.pdf"
        out_path = out_dir / out_name
        build_pdf_from_sheet(sheet_name, rows, out_path)
        print(f"Feuille {i} '{sheet_name}' -> {out_path}")

    wb.close()
    print("Terminé.")


if __name__ == "__main__":
    main()
