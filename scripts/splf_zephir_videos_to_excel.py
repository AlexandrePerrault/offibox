#!/usr/bin/env python3
"""
Récupère les URLs des vidéos du Guide ZÉPHIR (SPLF) et génère un Excel.
Col A = Produit, Col B = CIP (à compléter), Col C = URL.
Usage: python splf_zephir_videos_to_excel.py [--output fichier.xlsx]
"""
import re
import sys
from pathlib import Path

try:
    import requests
    from bs4 import BeautifulSoup
    import openpyxl
    from openpyxl.styles import Font, Alignment
except ImportError:
    print("Dépendances: requests, beautifulsoup4, openpyxl")
    print("  python -m pip install requests beautifulsoup4 openpyxl")
    sys.exit(1)

URL = "https://splf.fr/videos-zephir/"
OUTPUT_DEFAULT = "splf_zephir_videos.xlsx"

# Correspondance slug URL -> libellé produit (Guide ZÉPHIR SPLF)
SLUG_TO_PRODUCT = {
    "chambreinhalation": "Aérosol Doseur avec Chambre d'Inhalation",
    "airomir": "AIROMIR® AUTOHALER® (Salbutamol)",
    "alvesco": "ALVESCO® (Ciclésonide)",
    "anoro": "ANORO® ELLIPTA® (Umeclidinium + Vilanterol)",
    "asmanex": "ASMANEX® TWISTHALER® (Fuorate de mométasone)",
    "asmelor": "ASMELOR® NOVOLIZER® (Formotérol)",
    "atectura": "ATECTURA® BREEZHALER® (Indacatérol + mometasone)",
    "atrovent": "ATROVENT® (Ipratropium bromure)",
    "beclojet": "BECLOJET® (Béclométasone)",
    "beclometasone": "BECLOMETASONE TEVA® (Béclométasone)",
    "beclospray": "BECLOSPRAY® (Béclométasone)",
    "becotide": "BECOTIDE® (Béclométasone)",
    "bemedrex": "BEMEDREX® EASYHALER® (Béclométasone)",
    "bricanyl": "BRICANYL® TURBUHALER® (Terbutaline)",
    "bronchidual": "BRONCHODUAL® (Fénotérol + Ipratropium bromure)",
    "bronchodual": "BRONCHODUAL® (Fénotérol + Ipratropium bromure)",
    "duoresp": "DUORESP® SPIROMAX® (Budésonide + Formotérol)",
    "ecobec": "ECOBEC® (Béclométasone)",
    "elebrato": "ELEBRATO® ELLIPTA® (Fluticasone + Uméclidinium + Vilanterol)",
    "enerzair": "ENERZAIR® BREEZHALER® (Indacatérol + glycopyrronium + mometasone)",
    "flixotide": "FLIXOTIDE® (Fluticasone)",
    "flixotidediskus": "FLIXOTIDE® DISKUS® (Fluticasone)",
    "flutiform": "FLUTIFORM® (Fluticasone + Formotérol)",
    "foradil": "FORADIL® Aerolizer® (Formotérol)",
    "formoair": "FORMOAIR® (Formotérol)",
    "formodual": "FORMODUAL NEXThaler® (Béclométasone + Formotérol)",
    "formoterolbiogaran": "FORMOTEROL BIOGARAN® (Formotérol)",
    "formoteroleg": "FORMOTEROL EG® (Formotérol)",
    "formoterolmylan": "FORMOTEROL MYLAN® (Formotérol)",
    "formoterolzentiva": "FORMOTEROL ZENTIVA® (Formotérol)",
    "forspiro": "FORSPIRO® (Budésonide + Formotérol)",
    "gibiter": "GIBITER® Easyhaler® (Budésonide + Formotérol)",
    "incruse": "INCRUSE® ELLIPTA® (Bromure d'uméclidinium)",
    "innovair": "INNOVAIR NEXThaler® (Béclométasone + Formotérol)",
    "innovairnexthaler": "INNOVAIR NEXThaler® (Béclométasone + Formotérol)",
    "laventair": "LAVENTAIR® ELLIPTA® (Uméclidinium + Vilantérol)",
    "miflasone": "MIFLASONE® Aerolizer® (Béclométasone)",
    "novopulmon": "NOVOPULMON® NOVOLIZER® (Budésonide)",
    "onbrez": "ONBREZ® BREEZHALER® (Indacatérol)",
    "propionatefluticasonesalmeterolbiogaran": "Propionate de fluticasone/salmétérol BIOGARAN®",
    "propionatefluticasonesalmeterolmylan": "Propionate de fluticasone/salmétérol MYLAN",
    "pulmicort": "PULMICORT® TURBUHALER® (Budésonide)",
    "qvarspray": "QVARSPRAY® (Béclométasone)",
    "qvar": "QVAR® AUTOHALER® (Béclométasone)",
    "relvar": "RELVAR® ELLIPTA® (Fluticasone + Vilantérol)",
    "revinty": "REVINTY® ELLIPTA® (Fluticasone + Vilantérol)",
    "salbutamolteva": "SALBUTAMOL TEVA® (Salbutamol)",
    "seebri": "SEEBRI® BREEZHALER® (Bromure de glycopyrronium)",
    "seretidediskus": "SERETIDE® DISKUS® (Fluticasone + Salmétérol)",
    "seretide": "SERETIDE® SPRAY® (Fluticasone + Salmétérol)",
    "serevent": "SEREVENT® (Salmétérol)",
    "sereventdiskus": "SEREVENT® DISKUS® (Salmétérol)",
    "spiolto": "SPIOLTO® RESPIMAT® (Tiotropium + Olodatérol)",
    "spiriva": "SPIRIVA® HANDIHALER® (Tiotropium bromure)",
    "spirivarespimat": "SPIRIVA® RESPIMAT® (Tiotropium bromure)",
    "striverdi": "STRIVERDI® RESPIMAT® (Olodatérol)",
    "symbicort": "SYMBICORT® TURBUHALER® (Budésonide + Formotérol)",
    "symbicortrapihaler": "SYMBICORT® RAPIHALER® (Budésonide + Formotérol)",
    "tiotropiumbiogaran": "TIOTROPIUM-BIOGARAN (Tiotropium)",
    "tiotropiumviatris": "TIOTROPIUM-VIATRIS (Tiotropium)",
    "trelegy": "TRELEGY® ELLIPTA® (Fluticasone + Uméclidinium + Vilantérol)",
    "trimbow": "TRIMBOW NEXThaler® (Béclométasone + Formotérol + Glycopyrronium)",
    "trimbownexthaler": "TRIMBOW NEXThaler® (Béclométasone + Formotérol + Glycopyrronium)",
    "trixeo": "TRIXEO AEROSPHERE® (Formotérol + glycopyrronium + budesonide)",
    "ultibro": "ULTIBRO® BREEZHALER® (Indacatérol + Glycopyrronium bromure)",
    "ventilastin": "VENTILASTIN® NOVOLIZER® (Salbutamol)",
    "ventoline": "VENTOLINE® (Salbutamol)",
}


def fetch_page():
    r = requests.get(URL, timeout=30, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"})
    r.raise_for_status()
    r.encoding = r.apparent_encoding or "utf-8"
    return r.text


def extract_entries(html):
    soup = BeautifulSoup(html, "html.parser")
    entries = []

    # La page liste des blocs avec liens vidéo. Chercher les liens dans le contenu principal.
    # Structure typique: titre (h2/h3/h4) ou lien avec texte = nom du produit, href = page vidéo
    main = soup.find("main") or soup.find("article") or soup.find("div", class_=re.compile(r"content|entry|post", re.I)) or soup.body
    if not main:
        main = soup

    # Liens qui pointent vers des pages (souvent /videos-zephir/xxx ou similaires)
    seen_hrefs = set()
    for a in main.find_all("a", href=True):
        href = a.get("href", "").strip()
        if not href or href.startswith("#") or href in seen_hrefs:
            continue
        # Lien absolu
        if href.startswith("/"):
            href = "https://splf.fr" + href
        elif not href.startswith("http"):
            continue
        if "splf.fr" not in href and "youtube" not in href and "vimeo" not in href:
            continue
        seen_hrefs.add(href)
        # Nom du produit: texte du lien, ou alt de l'image, ou titre, ou déduit du slug URL
        name = (a.get_text(strip=True) or "").strip()
        if not name and a.find("img"):
            name = a.find("img").get("alt") or ""
        if not name:
            name = a.get("title") or ""
        name = name.strip()
        if not name:
            # Déduire du slug: .../portfolio-2/anoro/ -> nom produit
            slug = href.rstrip("/").split("/")[-1] or ""
            slug_flat = slug.replace("-", "").replace("_", "").lower()
            if slug in ("portfolio-2", "") or not slug:
                continue
            name = SLUG_TO_PRODUCT.get(slug_flat) or SLUG_TO_PRODUCT.get(slug.lower()) or (
                slug.replace("-", " ").replace("_", " ").title()
            )
            if len(name) < 2:
                continue
        # Éviter les liens "Charger plus", "Rechercher", etc.
        if re.search(r"charger plus|rechercher|connexion|accueil|mentions", name, re.I):
            continue
        if len(name) < 2:
            continue
        entries.append({"produit": name, "url": href})

    # Alternative: chercher des titres (h2, h3) suivis ou contenant un lien
    for tag in main.find_all(["h2", "h3", "h4"]):
        text = tag.get_text(strip=True)
        a = tag.find("a", href=True)
        if a and text and len(text) > 2:
            href = a.get("href", "").strip()
            if href.startswith("/"):
                href = "https://splf.fr" + href
            if href.startswith("http") and ("splf" in href or "youtube" in href or "vimeo" in href):
                key = (text[:80], href)
                if not any(e["produit"] == text and e["url"] == href for e in entries):
                    entries.append({"produit": text, "url": href})

    # Dédupliquer par (produit, url)
    by_key = {}
    for e in entries:
        k = (e["produit"][:100], e["url"])
        if k not in by_key:
            by_key[k] = e
    return list(by_key.values())


def main():
    out_path = OUTPUT_DEFAULT
    if "--output" in sys.argv:
        i = sys.argv.index("--output")
        if i + 1 < len(sys.argv):
            out_path = sys.argv[i + 1]

    print("Récupération de la page SPLF...")
    try:
        html = fetch_page()
    except Exception as e:
        print(f"Erreur: {e}")
        sys.exit(1)

    print("Extraction des liens...")
    entries = extract_entries(html)
    if not entries:
        print("Aucun lien trouvé. La structure de la page a peut-être changé.")
        sys.exit(1)

    print(f"Nombre d'entrées: {len(entries)}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ZÉPHIR"
    ws.append(["Produit", "CIP", "URL"])
    for c in "ABC":
        ws[f"{c}1"].font = Font(bold=True)
        ws[f"{c}1"].alignment = Alignment(horizontal="center", wrap_text=True)
    for e in entries:
        ws.append([e["produit"], "", e["url"]])
    for col in ["A", "B", "C"]:
        ws.column_dimensions[col].width = 50 if col == "A" else (18 if col == "B" else 70)

    wb.save(out_path)
    print(f"Fichier enregistré: {out_path}")
    print("Colonne CIP (B): à compléter avec la base des médicaments (BDM) si besoin.")


if __name__ == "__main__":
    main()
